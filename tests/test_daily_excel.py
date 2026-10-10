"""Kunlik Excel hisobot generatori (ma'lumotlar bazasiz, faqat tayyor lug'atlar bilan)."""
import io
from datetime import date, datetime

from openpyxl import load_workbook

from app.services.daily_excel import build_workbook


def _sample() -> dict:
    return {
        "date": date(2026, 10, 3),
        "generated_at": datetime(2026, 10, 4, 18, 42),
        "production": [
            {"id": 11, "machine_name": "Stanok-1", "note": None, "items": [
                {"product_name": "Paket", "unit": "dona", "partner_name": "Hamkor A", "quantity": 1000},
                {"product_name": "Plyonka", "unit": "kg", "partner_name": "Hamkor B", "quantity": 250.5},
            ]},
            {"id": 12, "machine_name": "Stanok-2", "note": "=1+1", "items": [
                {"product_name": "Paket", "unit": "dona", "partner_name": "Hamkor A", "quantity": 500},
            ]},
        ],
        "receipts": [
            {"partner_name": "Hamkor A", "brand": "PE-100", "unit": "kg", "quantity": 50, "unit_price": 12000.5,
             "supplier_name": "Ta'minotchi", "payment_type": "naqd", "note": "birinchi partiya"},
            {"partner_name": "Hamkor B", "brand": "PP-7", "unit": "kg", "quantity": 20, "unit_price": None,
             "supplier_name": None, "payment_type": None, "note": None},
        ],
        "sales": [
            {"id": 7, "customer_name": "Mijoz 1", "status": "paid", "debt": 0, "note": None, "items": [
                {"product_name": "Paket", "unit": "dona", "partner_name": None,
                 "quantity": 300, "unit_price": 1500, "total_price": 450000},
            ]},
            {"id": 8, "customer_name": "Mijoz 2", "status": "partial", "debt": 24999.97, "note": "2 oyga", "items": [
                {"product_name": "Plyonka", "unit": "kg", "partner_name": "Hamkor B",
                 "quantity": 10.5, "unit_price": 3333.33, "total_price": 34999.97},
            ]},
        ],
        "payments": [
            {"sale_id": 7, "customer_name": "Mijoz 1", "amount": 450000, "note": "Sotuv vaqtida"},
            {"sale_id": 8, "customer_name": "Mijoz 2", "amount": 10000, "note": "Sotuv vaqtida"},
            {"sale_id": 3, "customer_name": "Mijoz 3", "amount": 5000, "note": "Eski qarzdan"},
        ],
        "expenses": [
            {"id": 1, "category": "Oshxona", "description": "Kartoshka", "amount": 30000},
            {"id": 2, "category": "Elektr", "description": "Sentabr", "amount": 100000},
            {"id": 3, "category": "Oshxona", "description": None, "amount": 220000.5},
        ],
    }


def _empty() -> dict:
    d = _sample()
    d.update(production=[], receipts=[], sales=[], payments=[], expenses=[])
    return d


def _load(data, values=True):
    return load_workbook(io.BytesIO(build_workbook(data)), data_only=values)


def _find(ws, text):
    """Matn yozilgan birinchi katakni topadi -> (qator, ustun)."""
    for row in ws.iter_rows():
        for c in row:
            if c.value == text:
                return c.row, c.column
    raise AssertionError(f"'{text}' topilmadi")


def test_sheet_names_and_valid_xlsx():
    assert _load(_sample()).sheetnames == ["Ishlab chiqarish", "Moliya"]


def test_production_rows_and_totals():
    ws = _load(_sample())["Ishlab chiqarish"]
    r, _ = _find(ws, "Jami")
    assert ws.cell(r, 5).value == 1750.5  # 1000 + 250.5 + 500

    rp, _ = _find(ws, "Mahsulot bo'yicha jami")
    assert {ws.cell(rp + i, 4).value: ws.cell(rp + i, 5).value for i in (1, 2)} == {"Paket": 1500, "Plyonka": 250.5}

    rm, _ = _find(ws, "Stanok bo'yicha jami")
    assert {ws.cell(rm + i, 4).value: ws.cell(rm + i, 5).value for i in (1, 2)} == {"Stanok-1": 1250.5, "Stanok-2": 500}


def test_totals_are_formulas_with_cached_values():
    formulas = _load(_sample(), values=False)["Ishlab chiqarish"]
    r, _ = _find(formulas, "Jami")
    assert str(formulas.cell(r, 5).value).startswith("=SUM(")
    assert _load(_sample())["Ishlab chiqarish"].cell(r, 5).value == 1750.5


def test_text_starting_with_equals_is_not_a_formula():
    ws = _load(_sample(), values=False)["Ishlab chiqarish"]
    cell = next(c for row in ws.iter_rows() for c in row if c.value == "=1+1")
    assert cell.data_type == "s"


def test_finance_totals_and_net():
    ws = _load(_sample())["Moliya"]
    assert ws.cell(_find(ws, "Jami xomashyo xaridi")[0], 7).value == 600025.0  # 50 * 12000.50
    assert ws.cell(_find(ws, "Jami sotuv")[0], 7).value == 484999.97  # 450000 + 34999.97
    assert ws.cell(_find(ws, "Jami qabul qilingan to'lov")[0], 7).value == 465000
    assert ws.cell(_find(ws, "Jami xarajat")[0], 7).value == 350000.5
    net, _ = _find(ws, "Sof natija (sotuv − xarajat − xomashyo xaridi)")
    assert ws.cell(net, 7).value == -465025.53  # 484999.97 - 350000.5 - 600025


def test_sale_status_and_day_debt():
    ws = _load(_sample())["Moliya"]
    r, _ = _find(ws, "Qisman")
    assert ws.cell(r, 2).value == "Mijoz 2"
    paid_row, _ = _find(ws, "To'langan")
    assert ws.cell(paid_row, 2).value == "Mijoz 1"
    debt, _ = _find(ws, "Shu kungi sotuvlardan hozirgi qarz")
    assert ws.cell(debt, 7).value == 24999.97


def test_expenses_grouped_by_category():
    ws = _load(_sample())["Moliya"]
    r, _ = _find(ws, "Kategoriya bo'yicha jami")
    assert {ws.cell(r + i, 2).value: ws.cell(r + i, 7).value for i in (1, 2)} == {"Elektr": 100000, "Oshxona": 250000.5}


def test_empty_day_still_builds():
    wb = _load(_empty())
    assert wb.sheetnames == ["Ishlab chiqarish", "Moliya"]
    _find(wb["Ishlab chiqarish"], "Bu kunda ishlab chiqarish hisoboti kiritilmagan")
    _find(wb["Moliya"], "Bu kunda sotuv yo'q")
    _find(wb["Moliya"], "Bu kunda to'lov qabul qilinmagan")
    net, _ = _find(wb["Moliya"], "Sof natija (sotuv − xarajat − xomashyo xaridi)")
    assert wb["Moliya"].cell(net, 7).value == 0
