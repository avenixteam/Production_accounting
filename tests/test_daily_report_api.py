"""Kunlik Excel hisobot endpointi (butun tizim orqali). `pytest` bilan ishga tushiring."""
import io
from datetime import date, timedelta

from openpyxl import load_workbook
from test_flow import mk, sale_body, setup_master


def _find(ws, text):
    for row in ws.iter_rows():
        for c in row:
            if c.value == text:
                return c.row
    raise AssertionError(f"'{text}' topilmadi")


def _get(client, d):
    r = client.get("/reports/daily/excel", params={"date": d.isoformat()})
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert f"kunlik_hisobot_{d.isoformat()}.xlsx" in r.headers["content-disposition"]
    return load_workbook(io.BytesIO(r.content), data_only=True)


def test_backdated_entries_land_in_their_own_day(client):
    partner, rm, product, machine, customer = setup_master(client)
    yesterday = date.today() - timedelta(days=1)
    y = yesterday.isoformat()

    mk(client, "/receipts", {"partner_id": partner["id"], "raw_material_id": rm["id"],
                             "quantity": 100, "unit_price": 1000, "date": y})
    mk(client, "/production", {"date": y, "machine_id": machine["id"],
                               "items": [{"product_id": product["id"], "partner_id": partner["id"], "quantity": 80}]})
    mk(client, "/sales", sale_body(customer, product, 30, 2000, "partial", paid_amount=20000, date=y))
    mk(client, "/expenses/categories", {"name": "Elektr"})
    mk(client, "/expenses", {"category": "Elektr", "amount": 5000, "date": y})

    wb = _get(client, yesterday)
    assert wb.sheetnames == ["Ishlab chiqarish", "Moliya"]
    ws = wb["Ishlab chiqarish"]
    assert ws.cell(_find(ws, "Jami"), 5).value == 80
    fin = wb["Moliya"]
    assert fin.cell(_find(fin, "Jami sotuv"), 7).value == 60000
    assert fin.cell(_find(fin, "Jami xarajat"), 7).value == 5000
    assert fin.cell(_find(fin, "Jami xomashyo xaridi"), 7).value == 100000
    assert fin.cell(_find(fin, "Jami qabul qilingan to'lov"), 7).value == 20000
    assert fin.cell(_find(fin, "Shu kungi sotuvlardan hozirgi qarz"), 7).value == 40000

    today_wb = _get(client, date.today())
    assert "kiritilmagan" in today_wb["Ishlab chiqarish"]["A5"].value
    assert today_wb["Moliya"].cell(_find(today_wb["Moliya"], "Jami sotuv"), 7).value == 0


def test_default_date_is_today_and_empty_day_works(client):
    r = client.get("/reports/daily/excel")
    assert r.status_code == 200
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Ishlab chiqarish", "Moliya"]


def test_invalid_date_rejected(client):
    assert client.get("/reports/daily/excel", params={"date": "kecha"}).status_code == 422
